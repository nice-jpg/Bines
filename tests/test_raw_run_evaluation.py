from __future__ import annotations

import importlib
import sys
from pathlib import Path
from types import SimpleNamespace

import pytest
from langchain_core.messages import AIMessage, HumanMessage, ToolMessage
from langchain_core.messages.tool import tool_call
from openpyxl import Workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))


def _raw_run():
    return importlib.import_module("raw_run")


INDEX_HEADERS = [
    "merchant name",
    "distance",
    "rating",
    "total product count",
    "total review count",
]
PRODUCT_HEADERS = [
    "product name",
    "price",
    "original price",
    "discount price",
    "monthly sales",
]


def _write_workbook(
    path: Path,
    merchants: list[tuple[str, str, list[str]]],
) -> None:
    workbook = Workbook()
    index = workbook.active
    index.title = "Sheet1"
    index.append(INDEX_HEADERS)
    for merchant_name, sheet_name, products in merchants:
        index.append([merchant_name, "1km", 4.8, len(products), 100])
        worksheet = workbook.create_sheet(sheet_name)
        worksheet.append(PRODUCT_HEADERS)
        for product in products:
            worksheet.append([product, 10, 12, 10, 30])
    workbook.save(path)
    workbook.close()


def _ai_call(*calls) -> AIMessage:
    return AIMessage(content="", tool_calls=list(calls))


def test_eval_combines_quantity_context_and_correctness_grades(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    raw_run = _raw_run()
    result_path = tmp_path / "result.xlsx"
    truth_path = tmp_path / "truth.xlsx"
    _write_workbook(
        result_path,
        [
            (" A店 ", "A店", [" 商品1 ", "商品1", "商品2", "商品3", ""]),
            ("B店", "B店", ["商品4"]),
            ("result-only", "result-only", ["不参与正确性比较"]),
        ],
    )
    truth_workbook = Workbook()
    truth_workbook.remove(truth_workbook.active)
    for merchant_name, products in [
        ("A店", ["商品1", "真值商品"]),
        ("B店", ["商品4"]),
        ("truth-only", ["不参与比较"]),
    ]:
        worksheet = truth_workbook.create_sheet(merchant_name)
        worksheet.append(["商品名", "售价", "销量"])
        for product in products:
            worksheet.append([product, 10, 30])
    truth_workbook.save(truth_path)
    truth_workbook.close()
    messages = [
        HumanMessage(content="执行应用探测任务"),
        _ai_call(
            tool_call(
                name="notify_user",
                args={"current_path": "meituan/美食"},
                id="notice-list",
            ),
            tool_call(name="run_package", args={}, id="run-package"),
        ),
        ToolMessage(content="ok", tool_call_id="run-package"),
        _ai_call(
            tool_call(
                name="notify_user",
                args={"current_path": "meituan/美食/商家"},
                id="notice-merchant",
            ),
            tool_call(name="uiautomate", args={}, id="ui"),
            tool_call(name="query_manual", args={}, id="manual"),
        ),
        ToolMessage(content="xml", tool_call_id="ui"),
        _ai_call(tool_call(name="think", args={"thought": "inspect"}, id="think")),
        _ai_call(
            tool_call(
                name="notify_user",
                args={"current_path": "meituan/美食"},
                id="notice-back",
            ),
            tool_call(name="swipe_up", args={}, id="swipe"),
        ),
    ]
    result = SimpleNamespace(state={"messages": messages})
    monkeypatch.setattr(raw_run, "RESULT_WORKBOOK_PATH", result_path)
    monkeypatch.setattr(raw_run, "TRUTH_WORKBOOK_PATH", truth_path)

    # The result-only merchant is counted but is not part of correctness comparison.
    assert raw_run.eval(result) == {
        "total": 46,
        "dimensions": {
            "result_count_grade": 8,
            "context_grade": 40,
            "correctness_grade": -2,
        },
        "merchant_count": 3,
        "product_count": 5,
        "incorrect_product_count": 2,
    }


def test_context_grade_returns_zero_without_operations() -> None:
    raw_run = _raw_run()
    result = SimpleNamespace(
        state={
            "messages": [
                HumanMessage(content="start"),
                _ai_call(
                    tool_call(
                        name="notify_user",
                        args={"current_path": "meituan/美食/商家"},
                        id="notice",
                    )
                ),
            ]
        }
    )

    assert raw_run._context_grade(result) == 0


def test_workbook_reads_plain_merchant_name_without_hyperlink(tmp_path: Path) -> None:
    raw_run = _raw_run()
    path = tmp_path / "valid.xlsx"
    _write_workbook(path, [("A店", "A店", ["商品1"])])

    assert raw_run._read_merchant_products(path) == {"A店": {"商品1"}}


def test_workbook_requires_matching_merchant_worksheet_name(tmp_path: Path) -> None:
    raw_run = _raw_run()
    path = tmp_path / "invalid.xlsx"
    _write_workbook(path, [("A店", "A店", ["商品1"])])
    workbook = raw_run.load_workbook(path)
    workbook["A店"].title = "其他名称"
    workbook.save(path)
    workbook.close()

    with pytest.raises(ValueError, match="same name"):
        raw_run._read_merchant_products(path)


def test_workbook_requires_product_name_header(tmp_path: Path) -> None:
    raw_run = _raw_run()
    path = tmp_path / "invalid.xlsx"
    _write_workbook(path, [("A店", "A店", ["商品1"])])
    workbook = raw_run.load_workbook(path)
    workbook["A店"]["A1"] = "name"
    workbook.save(path)
    workbook.close()

    with pytest.raises(ValueError, match="product name"):
        raw_run._read_merchant_products(path)


def test_workbook_matches_plain_merchant_name_with_apostrophe(tmp_path: Path) -> None:
    raw_run = _raw_run()
    path = tmp_path / "apostrophe.xlsx"
    merchant_name = "A 店's products"
    _write_workbook(path, [(merchant_name, merchant_name, [" 商品1 ", "商品1"])])

    assert raw_run._read_merchant_products(path) == {merchant_name: {"商品1"}}


def test_truth_reader_accepts_new_indexed_workbook(tmp_path: Path) -> None:
    raw_run = _raw_run()
    path = tmp_path / "truth.xlsx"
    _write_workbook(path, [("A店", "A店", ["商品1"])])

    assert raw_run._read_truth_merchant_products(path) == {"A店": {"商品1"}}


def test_workbook_wraps_malformed_xlsx_error(tmp_path: Path) -> None:
    raw_run = _raw_run()
    path = tmp_path / "broken.xlsx"
    path.write_bytes(b"not an xlsx archive")

    with pytest.raises(ValueError, match="Unable to read evaluation workbook"):
        raw_run._read_merchant_products(path)


def test_eval_raises_when_truth_workbook_is_missing(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    raw_run = _raw_run()
    result_path = tmp_path / "result.xlsx"
    _write_workbook(result_path, [("A店", "A店", ["商品1"])])
    monkeypatch.setattr(raw_run, "RESULT_WORKBOOK_PATH", result_path)
    monkeypatch.setattr(raw_run, "TRUTH_WORKBOOK_PATH", tmp_path / "missing.xlsx")

    with pytest.raises(FileNotFoundError, match="missing.xlsx"):
        raw_run.eval(SimpleNamespace(state={"messages": []}))


@pytest.mark.parametrize(
    "state",
    [
        None,
        {},
        {"messages": "not-a-message-list"},
        {"messages": [{"role": "user", "content": "not BaseMessage"}]},
    ],
)
def test_context_grade_rejects_invalid_agent_state(state) -> None:
    raw_run = _raw_run()

    with pytest.raises(ValueError):
        raw_run._context_grade(SimpleNamespace(state=state))
