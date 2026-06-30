from __future__ import annotations

from concurrent.futures import ThreadPoolExecutor
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

from openpyxl import load_workbook
from openpyxl.workbook.workbook import Workbook as OpenpyxlWorkbook
from tools.content_provider import ContentProvider, _read_xlsx, create_content_provider_tools


class ContentProviderTests(unittest.TestCase):
    def test_create_append_and_update_excel_file_under_workspace(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            provider = ContentProvider(tmp_dir)
            created = Path(
                provider.create_excel_file(
                    "reports/stores.xlsx",
                    headers_json='["店名", "评分"]',
                    rows_json='[["A店", 4.8]]',
                )
            )
            provider.append_excel_rows("reports/stores.xlsx", '[["B店", 4.5]]')
            provider.update_excel_cell("reports/stores.xlsx", "C2", "已采集")

            self.assertEqual(created, Path(tmp_dir).resolve() / "reports/stores.xlsx")
            self.assertEqual(
                _read_xlsx(created),
                [
                    ["店名", "评分"],
                    ["A店", "4.8", "已采集"],
                    ["B店", "4.5"],
                ],
            )

    def test_rejects_paths_outside_workspace(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            provider = ContentProvider(tmp_dir)

            with self.assertRaises(ValueError):
                provider.create_excel_file("../outside.xlsx")

    def test_create_rename_and_write_specific_sheet(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            provider = ContentProvider(tmp_dir)
            path = Path(
                provider.create_excel_file(
                    "result.xlsx",
                    headers_json='["商品名", "价格"]',
                    rows_json='[["商品A", 10]]',
                    sheet_name="商家A完整名称",
                )
            )
            provider.create_excel_sheet(
                "result.xlsx",
                "临时名称",
                headers_json='["商品名", "价格"]',
            )
            provider.rename_excel_sheet("result.xlsx", "临时名称", "商家B完整名称")
            provider.write_excel_sheet(
                "result.xlsx",
                "商家B完整名称",
                headers_json='["商品名", "价格"]',
                rows_json='[["商品B", 20]]',
            )
            provider.append_excel_rows(
                "result.xlsx",
                '[["商品C", 30]]',
                sheet_name="商家B完整名称",
            )
            provider.update_excel_cell(
                "result.xlsx",
                "C2",
                "在售",
                sheet_name="商家B完整名称",
            )

            workbook = load_workbook(path, data_only=True)
            try:
                self.assertEqual(workbook.sheetnames, ["商家A完整名称", "商家B完整名称"])
            finally:
                workbook.close()
            self.assertEqual(
                _read_xlsx(path, "商家A完整名称"),
                [["商品名", "价格"], ["商品A", "10"]],
            )
            self.assertEqual(
                _read_xlsx(path, "商家B完整名称"),
                [["商品名", "价格"], ["商品B", "20", "在售"], ["商品C", "30"]],
            )

    def test_rejects_invalid_or_duplicate_sheet_names(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            provider = ContentProvider(tmp_dir)
            provider.create_excel_file("result.xlsx", sheet_name="商家A")

            with self.assertRaisesRegex(ValueError, "already exists"):
                provider.create_excel_sheet("result.xlsx", "商家a")
            with self.assertRaisesRegex(ValueError, "must not contain"):
                provider.create_excel_sheet("result.xlsx", "商家/A")
            with self.assertRaisesRegex(ValueError, "31 characters"):
                provider.create_excel_sheet("result.xlsx", "A" * 32)

    def test_content_provider_registers_multi_sheet_tools(self) -> None:
        tool_names = [tool.name for tool in create_content_provider_tools(ContentProvider())]

        self.assertIn("create_excel_sheet", tool_names)
        self.assertIn("rename_excel_sheet", tool_names)
        self.assertIn("write_excel_sheet", tool_names)

    def test_concurrent_provider_instances_serialize_updates_to_same_workbook(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            first_provider = ContentProvider(tmp_dir)
            second_provider = ContentProvider(tmp_dir)
            path = Path(
                first_provider.create_excel_file(
                    "result.xlsx",
                    headers_json='["商家名称"]',
                )
            )

            def append_row(index: int) -> None:
                provider = first_provider if index % 2 == 0 else second_provider
                provider.append_excel_rows("result.xlsx", f'[["商家{index}"]]')

            with ThreadPoolExecutor(max_workers=8) as executor:
                list(executor.map(append_row, range(40)))

            rows = _read_xlsx(path)
            self.assertEqual(rows[0], ["商家名称"])
            self.assertEqual(len(rows), 41)
            self.assertEqual({row[0] for row in rows[1:]}, {f"商家{i}" for i in range(40)})

    def test_failed_save_preserves_existing_workbook(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            provider = ContentProvider(tmp_dir)
            path = Path(
                provider.create_excel_file(
                    "result.xlsx",
                    headers_json='["商家名称"]',
                    rows_json='[["原始商家"]]',
                )
            )
            original_bytes = path.read_bytes()

            def fail_after_partial_write(_workbook, target_path) -> None:
                Path(target_path).write_bytes(b"incomplete zip")
                raise OSError("simulated save failure")

            with patch.object(
                OpenpyxlWorkbook,
                "save",
                autospec=True,
                side_effect=fail_after_partial_write,
            ):
                with self.assertRaisesRegex(OSError, "simulated save failure"):
                    provider.append_excel_rows("result.xlsx", '[["不应写入"]]')

            self.assertEqual(path.read_bytes(), original_bytes)
            self.assertEqual(_read_xlsx(path), [["商家名称"], ["原始商家"]])
            self.assertEqual(list(path.parent.glob(".result.*.tmp.xlsx")), [])

    def test_corrupted_workbook_returns_recovery_guidance(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            provider = ContentProvider(tmp_dir)
            Path(tmp_dir, "result.xlsx").write_bytes(b"incomplete zip")

            with self.assertRaisesRegex(
                ValueError,
                "corrupted.*recreate it with create_excel_file",
            ):
                provider.append_excel_rows("result.xlsx", '[["商家A"]]')


if __name__ == "__main__":
    unittest.main()
