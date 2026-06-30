from __future__ import annotations

from collections.abc import Callable, Mapping, Sequence
from pathlib import Path
from typing import Any
from zipfile import BadZipFile

from agent import AgentRunResult, AgentRuntime
from langchain_core.messages import AIMessage, BaseMessage
from model import build_codex_model, build_model
from openpyxl import load_workbook
from openpyxl.formula import Tokenizer
from openpyxl.utils.exceptions import InvalidFileException
from tools.common import CaptchaAuthenticationTool, OperationNoticeTool, create_common_tools


AgentRunLoop = Callable[[], AgentRunResult]
Evaluator = Callable[[AgentRunResult], Mapping[str, Any]]
WORKSPACE_DIR = Path(__file__).resolve().parents[1] / "workspace"
RESULT_WORKBOOK_PATH = WORKSPACE_DIR / "result.xlsx"
TRUTH_WORKBOOK_PATH = WORKSPACE_DIR / "truth.xlsx"
MERCHANT_INDEX_SHEET = "Sheet1"
MERCHANT_INDEX_HEADERS = (
    "merchant name",
    "distance",
    "rating",
    "total product count",
    "total review count",
)
MERCHANT_NAME_HEADER = "merchant name"
PRODUCT_DETAIL_HEADERS = (
    "product name",
    "price",
    "original price",
    "discount price",
)
PRODUCT_NAME_HEADER = "product name"
SALES_HEADERS = ("monthly sales", "sales")
TRUTH_PRODUCT_NAME_HEADERS = ("product name", "product_name", "商品名")
INCORRECT_PRODUCT_PENALTY = 10


class PlainNotifier:
    def set(self, _) -> None:
        pass

    def __call__(self, text: str) -> None:
        print(f"[NOTIFY] {text}")


def run() -> None:
    notifier = PlainNotifier()
    operation_notice = OperationNoticeTool(notifier=notifier)
    captcha_authentication = CaptchaAuthenticationTool(notifier=notifier)
    runtime = AgentRuntime(
        model=build_codex_model(),
        tools=create_common_tools(
            operation_notice_tool=operation_notice,
            captcha_authentication_tool=captcha_authentication,
        ),
        name="main",
    )

    r = runtime.run_turn(
        [{"role": "user", "content": "执行应用探测任务"}],
        session_id="main",
        max_iterations=1000,
    )

    print(r)


def eval(result: AgentRunResult) -> Mapping[str, Any]:
    """Return a mind_controller-compatible structured evaluation."""

    result_merchants = _read_merchant_products(RESULT_WORKBOOK_PATH)
    truth_merchants = _read_truth_merchant_products(TRUTH_WORKBOOK_PATH)

    merchant_count = len(result_merchants)
    product_count = sum(len(products) for products in result_merchants.values())
    result_count_grade = merchant_count + product_count
    context_grade = _context_grade(result)
    incorrect_product_count = _incorrect_product_count(result_merchants, truth_merchants)
    correctness_grade = -INCORRECT_PRODUCT_PENALTY * incorrect_product_count
    total = result_count_grade + context_grade + correctness_grade

    return {
        "total": total,
        "dimensions": {
            "result_count_grade": result_count_grade,
            "context_grade": context_grade,
            "correctness_grade": correctness_grade,
        },
        "merchant_count": merchant_count,
        "product_count": product_count,
        "incorrect_product_count": incorrect_product_count,
    }


def _read_merchant_products(path: Path) -> dict[str, set[str]]:
    if not path.is_file():
        raise FileNotFoundError(f"Evaluation workbook does not exist: {path}")

    try:
        workbook = load_workbook(path, data_only=False)
    except (BadZipFile, InvalidFileException, OSError, ValueError) as exc:
        raise ValueError(f"Unable to read evaluation workbook {path}: {exc}") from exc

    merchants: dict[str, set[str]] = {}
    try:
        if MERCHANT_INDEX_SHEET not in workbook.sheetnames:
            raise ValueError(
                f"Evaluation workbook {path} must contain merchant index worksheet "
                f"{MERCHANT_INDEX_SHEET!r}"
            )
        index_worksheet = workbook[MERCHANT_INDEX_SHEET]
        header_row_number, headers = _read_headers(
            index_worksheet,
            path,
            MERCHANT_INDEX_HEADERS,
        )
        merchant_index = headers.index(MERCHANT_NAME_HEADER)
        for row in index_worksheet.iter_rows(min_row=header_row_number + 1):
            merchant_name, destination = _merchant_link(row, merchant_index, path)
            if not merchant_name:
                continue
            detail_worksheet = _linked_worksheet(
                workbook,
                destination,
                merchant_name,
                path,
            )
            products = _read_product_worksheet(detail_worksheet, path)
            merchants.setdefault(merchant_name, set()).update(products)
    finally:
        workbook.close()
    return merchants


def _read_headers(
    worksheet: Any,
    workbook_path: Path,
    required_headers: Sequence[str],
) -> tuple[int, list[str]]:
    header_row = None
    for row in worksheet.iter_rows():
        if any(_normalized_name(cell.value) for cell in row):
            header_row = row
            break
    if header_row is None:
        raise ValueError(
            f"Worksheet {worksheet.title!r} in {workbook_path} is empty"
        )

    headers = [_normalized_name(cell.value) for cell in header_row]
    nonempty_headers = [header for header in headers if header]
    missing_headers = [header for header in required_headers if header not in headers]
    if missing_headers:
        raise ValueError(
            f"Worksheet {worksheet.title!r} in {workbook_path} "
            f"is missing required headers: {missing_headers!r}"
        )
    if len(set(nonempty_headers)) != len(nonempty_headers):
        raise ValueError(
            f"Worksheet {worksheet.title!r} in {workbook_path} contains duplicate headers"
        )
    return header_row[0].row, headers


def _linked_worksheet(
    workbook: Any,
    destination: str,
    merchant_name: str,
    workbook_path: Path,
) -> Any:
    if not destination:
        raise ValueError(
            f"Merchant {merchant_name!r} in worksheet {MERCHANT_INDEX_SHEET!r} "
            f"of {workbook_path} must link to its product worksheet"
        )

    worksheet_name = _hyperlink_worksheet_name(destination)
    if not worksheet_name or worksheet_name not in workbook.sheetnames:
        raise ValueError(
            f"Merchant {merchant_name!r} in {workbook_path} links to unknown "
            f"worksheet destination {destination!r}"
        )
    if worksheet_name == MERCHANT_INDEX_SHEET:
        raise ValueError(
            f"Merchant {merchant_name!r} in {workbook_path} links to the index worksheet"
        )
    return workbook[worksheet_name]


def _merchant_link(
    row: Sequence[Any],
    merchant_index: int,
    workbook_path: Path,
) -> tuple[str, str]:
    if merchant_index >= len(row):
        return "", ""
    merchant_cell = row[merchant_index]
    value = _normalized_name(merchant_cell.value)
    if merchant_cell.data_type == "f":
        parsed = _parse_hyperlink_formula(value)
        if parsed is None:
            raise ValueError(
                f"Merchant cell {merchant_cell.coordinate} in {workbook_path} "
                "must use a HYPERLINK formula"
            )
        return parsed

    merchant_name = value
    if not merchant_name:
        return "", ""
    hyperlink = merchant_cell.hyperlink
    if hyperlink is None:
        hyperlinks = [cell.hyperlink for cell in row if cell.hyperlink is not None]
        if len(hyperlinks) == 1:
            hyperlink = hyperlinks[0]
    destination = hyperlink.location or hyperlink.target or "" if hyperlink else ""
    return merchant_name, destination


def _parse_hyperlink_formula(formula: str) -> tuple[str, str] | None:
    tokens = Tokenizer(formula).items
    if not tokens or tokens[0].type != "FUNC":
        return None
    function_name = tokens[0].value.removesuffix("(").casefold()
    if function_name != "hyperlink":
        return None
    arguments = [
        token.value[1:-1].replace('""', '"')
        for token in tokens
        if token.type == "OPERAND" and token.subtype == "TEXT"
    ]
    if len(arguments) != 2:
        return None
    destination, merchant_name = arguments
    return _normalized_name(merchant_name), _normalized_name(destination)


def _hyperlink_worksheet_name(destination: str) -> str:
    location = _normalized_name(destination).lstrip("#")
    if "!" not in location:
        return ""
    worksheet_name = location.rsplit("!", 1)[0]
    if worksheet_name.startswith("'") and worksheet_name.endswith("'"):
        worksheet_name = worksheet_name[1:-1].replace("''", "'")
    return worksheet_name


def _read_product_worksheet(worksheet: Any, workbook_path: Path) -> set[str]:
    header_row_number, headers = _read_headers(
        worksheet,
        workbook_path,
        PRODUCT_DETAIL_HEADERS,
    )
    if not any(header in headers for header in SALES_HEADERS):
        raise ValueError(
            f"Worksheet {worksheet.title!r} in {workbook_path} is missing required "
            f"sales header; expected one of {SALES_HEADERS!r}"
        )
    product_index = headers.index(PRODUCT_NAME_HEADER)
    products: set[str] = set()
    for row in worksheet.iter_rows(
        min_row=header_row_number + 1,
        values_only=True,
    ):
        product_name = _row_name(row, product_index)
        if product_name:
            products.add(product_name)
    return products


def _read_truth_merchant_products(path: Path) -> dict[str, set[str]]:
    if not path.is_file():
        raise FileNotFoundError(f"Evaluation workbook does not exist: {path}")

    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except (BadZipFile, InvalidFileException, OSError, ValueError) as exc:
        raise ValueError(f"Unable to read evaluation workbook {path}: {exc}") from exc

    merchants: dict[str, set[str]] = {}
    try:
        if MERCHANT_INDEX_SHEET in workbook.sheetnames:
            index_headers = {
                _normalized_name(cell.value)
                for row in workbook[MERCHANT_INDEX_SHEET].iter_rows(max_row=1)
                for cell in row
            }
            if set(MERCHANT_INDEX_HEADERS) <= index_headers:
                workbook.close()
                return _read_merchant_products(path)
        for worksheet in workbook.worksheets:
            rows = iter(worksheet.iter_rows(values_only=True))
            header_row = next(
                (row for row in rows if any(_normalized_name(value) for value in row)),
                None,
            )
            if header_row is None:
                continue
            headers = [_normalized_name(value) for value in header_row]
            product_headers = [
                header for header in TRUTH_PRODUCT_NAME_HEADERS if header in headers
            ]
            if len(product_headers) != 1:
                raise ValueError(
                    f"Worksheet {worksheet.title!r} in {path} must contain exactly "
                    f"one product header from {TRUTH_PRODUCT_NAME_HEADERS!r}"
                )
            product_index = headers.index(product_headers[0])
            products = {
                product_name
                for row in rows
                if (product_name := _row_name(row, product_index))
            }
            merchant_name = _normalized_name(worksheet.title)
            merchants.setdefault(merchant_name, set()).update(products)
    finally:
        workbook.close()
    return merchants


def _row_name(row: Sequence[Any], index: int) -> str:
    return _normalized_name(row[index]) if index < len(row) else ""


def _normalized_name(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _context_grade(result: AgentRunResult) -> int:
    state = getattr(result, "state", None)
    if not isinstance(state, Mapping):
        raise ValueError("AgentRunResult.state must be a mapping")
    messages = state.get("messages")
    if not isinstance(messages, Sequence) or isinstance(messages, (str, bytes)):
        raise ValueError("AgentRunResult.state['messages'] must be a sequence")
    if any(not isinstance(message, BaseMessage) for message in messages):
        raise ValueError("AgentRunResult.state['messages'] must contain BaseMessage values")

    current_path = ""
    operation_count = 0
    redundant_operation_count = 0
    for message in messages:
        if not isinstance(message, AIMessage):
            continue
        for tool_call in message.tool_calls:
            if not isinstance(tool_call, Mapping):
                raise ValueError("AIMessage.tool_calls must contain mapping values")
            tool_name = str(tool_call.get("name") or "")
            args = tool_call.get("args") or {}
            if tool_name == "notify_user":
                if not isinstance(args, Mapping):
                    raise ValueError("notify_user tool arguments must be a mapping")
                current_path = _normalized_name(args.get("current_path"))
                continue
            operation_count += 1
            if _is_merchant_path(current_path):
                redundant_operation_count += 1

    if operation_count == 0:
        return 0
    return round(100 * (1 - redundant_operation_count / operation_count))


def _is_merchant_path(path: str) -> bool:
    normalized = path.strip().rstrip("/")
    return normalized == "商家" or normalized.endswith("/商家")


def _incorrect_product_count(
    result_merchants: Mapping[str, set[str]],
    truth_merchants: Mapping[str, set[str]],
) -> int:
    shared_merchants = result_merchants.keys() & truth_merchants.keys()
    return sum(
        len(result_merchants[merchant] - truth_merchants[merchant])
        for merchant in shared_merchants
    )
