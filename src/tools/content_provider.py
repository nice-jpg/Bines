"""Workspace-scoped Excel output tools."""

from __future__ import annotations

import json
from pathlib import Path
import posixpath
import re

from langchain_core.tools import StructuredTool
from openpyxl import Workbook, load_workbook
from openpyxl.workbook.workbook import Workbook as OpenpyxlWorkbook
from openpyxl.worksheet.worksheet import Worksheet

WORKSPACE_ROOT = Path(__file__).resolve().parents[2] / "workspace"
SHEET_NAME = "Sheet1"
INVALID_SHEET_NAME_RE = re.compile(r"[\[\]:*?/\\]")


class ContentProvider:
    """Create and update Excel files under the workspace directory."""

    def __init__(self, workspace_root: str | Path = WORKSPACE_ROOT) -> None:
        self.workspace_root = Path(workspace_root).resolve()
        self.workspace_root.mkdir(parents=True, exist_ok=True)

    def create_excel_file(
        self,
        relative_path: str,
        headers_json: str = "[]",
        rows_json: str = "[]",
        sheet_name: str = SHEET_NAME,
    ) -> str:
        path = self._resolve_excel_path(relative_path)
        validated_name = _validate_sheet_name(sheet_name)
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = validated_name
        _replace_sheet_rows(worksheet, _table_from_json(headers_json, rows_json))
        _save_workbook(workbook, path)
        return str(path)

    def create_excel_sheet(
        self,
        relative_path: str,
        sheet_name: str,
        headers_json: str = "[]",
        rows_json: str = "[]",
    ) -> str:
        path = self._resolve_excel_path(relative_path)
        workbook = _load_existing_workbook(path)
        try:
            validated_name = _validate_sheet_name(sheet_name)
            _ensure_unique_sheet_name(workbook, validated_name)
            worksheet = workbook.create_sheet(validated_name)
            _replace_sheet_rows(worksheet, _table_from_json(headers_json, rows_json))
            _save_workbook(workbook, path)
        except Exception:
            workbook.close()
            raise
        return str(path)

    def rename_excel_sheet(self, relative_path: str, sheet_name: str, new_sheet_name: str) -> str:
        path = self._resolve_excel_path(relative_path)
        workbook = _load_existing_workbook(path)
        try:
            worksheet = _get_sheet(workbook, sheet_name)
            validated_name = _validate_sheet_name(new_sheet_name)
            if validated_name.casefold() != worksheet.title.casefold():
                _ensure_unique_sheet_name(workbook, validated_name)
            worksheet.title = validated_name
            _save_workbook(workbook, path)
        except Exception:
            workbook.close()
            raise
        return str(path)

    def write_excel_sheet(
        self,
        relative_path: str,
        sheet_name: str,
        headers_json: str = "[]",
        rows_json: str = "[]",
    ) -> str:
        path = self._resolve_excel_path(relative_path)
        workbook = _load_existing_workbook(path)
        try:
            worksheet = _get_sheet(workbook, sheet_name)
            _replace_sheet_rows(worksheet, _table_from_json(headers_json, rows_json))
            _save_workbook(workbook, path)
        except Exception:
            workbook.close()
            raise
        return str(path)

    def append_excel_rows(
        self,
        relative_path: str,
        rows_json: str,
        sheet_name: str = SHEET_NAME,
    ) -> str:
        path = self._resolve_excel_path(relative_path)
        if not path.exists():
            self.create_excel_file(relative_path, sheet_name=sheet_name)
        workbook = _load_existing_workbook(path)
        try:
            worksheet = _get_sheet(workbook, sheet_name)
            for row in _parse_json_rows(rows_json, allow_flat=False):
                worksheet.append(row)
            _save_workbook(workbook, path)
        except Exception:
            workbook.close()
            raise
        return str(path)

    def update_excel_cell(
        self,
        relative_path: str,
        cell: str,
        value: str,
        sheet_name: str = SHEET_NAME,
    ) -> str:
        path = self._resolve_excel_path(relative_path)
        if not path.exists():
            self.create_excel_file(relative_path, sheet_name=sheet_name)
        workbook = _load_existing_workbook(path)
        try:
            worksheet = _get_sheet(workbook, sheet_name)
            worksheet[cell] = value
            _save_workbook(workbook, path)
        except Exception:
            workbook.close()
            raise
        return str(path)

    def _resolve_excel_path(self, relative_path: str) -> Path:
        normalized = posixpath.normpath(str(relative_path or "").replace("\\", "/")).lstrip("/")
        if normalized in {"", "."}:
            raise ValueError("relative_path is required")
        path = (self.workspace_root / normalized).resolve()
        if self.workspace_root != path and self.workspace_root not in path.parents:
            raise ValueError("Excel path must stay inside workspace")
        if path.suffix.lower() != ".xlsx":
            raise ValueError("Excel path must end with .xlsx")
        path.parent.mkdir(parents=True, exist_ok=True)
        return path


def create_content_provider_tools(provider: ContentProvider | None = None) -> list[StructuredTool]:
    content_provider = provider or ContentProvider()
    return [
        StructuredTool.from_function(
            func=content_provider.create_excel_file,
            name="create_excel_file",
            description=(
                "Create an .xlsx file under workspace. Inputs: relative_path, optional "
                "headers_json, optional rows_json, and optional sheet_name."
            ),
        ),
        StructuredTool.from_function(
            func=content_provider.create_excel_sheet,
            name="create_excel_sheet",
            description=(
                "Create a worksheet in an existing .xlsx file. Inputs: relative_path, "
                "sheet_name, optional headers_json, and optional rows_json."
            ),
        ),
        StructuredTool.from_function(
            func=content_provider.rename_excel_sheet,
            name="rename_excel_sheet",
            description=(
                "Rename a worksheet. Inputs: relative_path, sheet_name, and new_sheet_name."
            ),
        ),
        StructuredTool.from_function(
            func=content_provider.write_excel_sheet,
            name="write_excel_sheet",
            description=(
                "Replace all rows in a specific worksheet. Inputs: relative_path, sheet_name, "
                "optional headers_json, and optional rows_json."
            ),
        ),
        StructuredTool.from_function(
            func=content_provider.append_excel_rows,
            name="append_excel_rows",
            description=(
                "Append rows to a specific worksheet. Inputs: relative_path, rows_json, "
                "and optional sheet_name."
            ),
        ),
        StructuredTool.from_function(
            func=content_provider.update_excel_cell,
            name="update_excel_cell",
            description=(
                "Update one cell in a specific worksheet. Inputs: relative_path, cell such "
                "as 'B2', value, and optional sheet_name."
            ),
        ),
    ]


def build_tools() -> list[StructuredTool]:
    return create_content_provider_tools()


def _parse_json_rows(value: str, allow_flat: bool) -> list[list[object]]:
    if not value:
        return []
    parsed = json.loads(value)
    if allow_flat and isinstance(parsed, list) and not _is_row_like(parsed):
        return [parsed]
    if isinstance(parsed, dict):
        return [[parsed[key] for key in parsed]]
    if not isinstance(parsed, list):
        raise ValueError("Expected JSON array")
    if not parsed:
        return []
    if allow_flat and not any(_is_row_like(item) for item in parsed):
        return [parsed]
    rows: list[list[object]] = []
    for item in parsed:
        if isinstance(item, dict):
            rows.append([item[key] for key in item])
        elif isinstance(item, list):
            rows.append(item)
        else:
            rows.append([item])
    return rows


def _is_row_like(value: object) -> bool:
    return isinstance(value, (list, dict))


def _table_from_json(headers_json: str, rows_json: str) -> list[list[object]]:
    headers = _parse_json_rows(headers_json, allow_flat=True)
    rows = _parse_json_rows(rows_json, allow_flat=False)
    table: list[list[object]] = []
    if headers:
        table.append(headers[0])
    table.extend(rows)
    return table


def _replace_sheet_rows(worksheet: Worksheet, rows: list[list[object]]) -> None:
    if worksheet.max_row:
        worksheet.delete_rows(1, worksheet.max_row)
    for row in rows:
        worksheet.append(row)


def _load_existing_workbook(path: Path) -> OpenpyxlWorkbook:
    if not path.exists():
        raise ValueError(f"Excel file does not exist: {path.name}")
    return load_workbook(path)


def _save_workbook(workbook: OpenpyxlWorkbook, path: Path) -> None:
    workbook.save(path)
    workbook.close()


def _get_sheet(workbook: OpenpyxlWorkbook, sheet_name: str) -> Worksheet:
    validated_name = _validate_sheet_name(sheet_name)
    if validated_name not in workbook.sheetnames:
        available = ", ".join(workbook.sheetnames) or "none"
        raise ValueError(f'Worksheet "{validated_name}" does not exist. Available sheets: {available}')
    return workbook[validated_name]


def _validate_sheet_name(sheet_name: str) -> str:
    name = str(sheet_name or "")
    if not name.strip():
        raise ValueError("sheet_name must not be empty")
    if len(name) > 31:
        raise ValueError("sheet_name must not exceed 31 characters")
    if INVALID_SHEET_NAME_RE.search(name):
        raise ValueError(r"sheet_name must not contain any of: []:*?/\\")
    return name


def _ensure_unique_sheet_name(workbook: OpenpyxlWorkbook, sheet_name: str) -> None:
    if any(existing.casefold() == sheet_name.casefold() for existing in workbook.sheetnames):
        raise ValueError(f'Worksheet "{sheet_name}" already exists')


def _read_xlsx(path: Path, sheet_name: str = SHEET_NAME) -> list[list[object]]:
    workbook = load_workbook(path, data_only=True)
    try:
        worksheet = _get_sheet(workbook, sheet_name)
        rows: list[list[object]] = []
        for row in worksheet.iter_rows(values_only=True):
            values = [_read_cell_value(value) for value in row]
            while values and values[-1] == "":
                values.pop()
            rows.append(values)
        while rows and not rows[-1]:
            rows.pop()
        return rows
    finally:
        workbook.close()


def _read_cell_value(value: object) -> object:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "1" if value else "0"
    return str(value)
